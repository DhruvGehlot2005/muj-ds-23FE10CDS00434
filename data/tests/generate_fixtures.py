"""
@file generate_fixtures.py
@description Generates anonymized test PDF and DOCX files for test suites
@module data.tests.generate_fixtures
"""

import os
import pymupdf
import docx

os.makedirs("data/tests", exist_ok=True)

# 1. Valid PDF Resume
doc_pdf = pymupdf.open()
page = doc_pdf.new_page()
resume_text = """Alex Morgan
Email: alex.morgan@testdomain.edu | Phone: +1-555-0144 | LinkedIn: linkedin.com/in/alexmorgan-dev
Summary: Dedicated Software Developer with solid foundation in algorithms, REST APIs, and full stack web development.

EDUCATION
Bachelor of Science in Computer Science | State University | GPA: 3.8 / 4.0 | 2020 - 2024

TECHNICAL SKILLS
Languages: Python, TypeScript, JavaScript, SQL, HTML, CSS
Frameworks & Tools: FastAPI, React, Node.js, Git, Docker, PostgreSQL

EXPERIENCE & PROJECTS
Software Engineering Intern | TechCorp Labs | Summer 2023
- Designed and implemented RESTful endpoints using Python and FastAPI, serving 10,000+ daily requests.
- Integrated automated testing pipelines with pytest, achieving 92% code coverage.
- Optimized database indexing in PostgreSQL, reducing query latency by 35%.

NLP Resume Parser Project | Academic Capstone | 2024
- Built a resume parsing pipeline using Python, PyMuPDF, and LLM APIs to extract structured candidate profiles.
- Implemented real-time ATS compatibility scoring and gap analysis.
"""
page.insert_text((50, 50), resume_text, fontsize=11)
pdf_path = os.path.join("data", "tests", "sample_resume.pdf")
doc_pdf.save(pdf_path)
doc_pdf.close()

# 2. Valid DOCX Resume
doc_docx = docx.Document()
doc_docx.add_heading("Taylor Reed", level=1)
doc_docx.add_paragraph("Email: taylor.reed@testdomain.edu | Phone: +1-555-0188 | GitHub: github.com/taylorreed")
doc_docx.add_heading("Education", level=2)
doc_docx.add_paragraph("B.S. in Data Science and Analytics, University of Technology, 2021 - 2025")
doc_docx.add_heading("Skills", level=2)
doc_docx.add_paragraph("Python, SQL, R, Pandas, NumPy, Scikit-learn, Tableau, Git")
doc_docx.add_heading("Experience", level=2)
doc_docx.add_paragraph("Data Analyst Intern at Insight Analytics. Developed automated reporting dashboards.")
docx_path = os.path.join("data", "tests", "sample_resume.docx")
doc_docx.save(docx_path)

# 3. Non-resume text PDF (>150 characters, no resume sections)
doc_recipe = pymupdf.open()
p2 = doc_recipe.new_page()
recipe_content = (
    "Chocolate Chip Cookies Recipe\n"
    "Ingredients: 2 cups all-purpose flour, 1 cup granulated sugar, 1 cup brown sugar, "
    "2 large eggs, 1 teaspoon baking soda, 1 cup chocolate chips, softened butter, vanilla extract.\n"
    "Instructions: Preheat oven to 375 degrees Fahrenheit. Whisk dry ingredients in a large bowl. "
    "Cream butter and sugars until fluffy. Fold in chocolate chips. Bake on parchment-lined sheet "
    "for 10 to 12 minutes until golden brown. Cool on wire rack before serving."
)
p2.insert_text((50, 50), recipe_content, fontsize=11)
recipe_path = os.path.join("data", "tests", "recipe_not_resume.pdf")
doc_recipe.save(recipe_path)
doc_recipe.close()

print("Successfully created test fixtures in data/tests!")
